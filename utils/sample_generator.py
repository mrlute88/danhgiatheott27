"""
sample_generator.py - Generates realistic standard Excel templates compliant with Circular 27/2020/TT-BGDĐT.
Used for instant testing and download by primary school teachers.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SAMPLE_STUDENTS = [
    {
        "stt": 1,
        "code": "HS0301",
        "name": "Nguyễn Hoàng Minh Khôi",
        "gender": "Nam",
        "dob": "12/03/2016",
        "score": 10.0,
        "level": "T",
        "note": "Tính nhẩm rất nhanh, chữ viết đẹp, tích cực phát biểu",
        "comment": ""
    },
    {
        "stt": 2,
        "code": "HS0302",
        "name": "Trần Ngọc Mai Anh",
        "gender": "Nữ",
        "dob": "25/08/2016",
        "score": 9.5,
        "level": "T",
        "note": "Nắm chắc kiến thức, cẩn thận, giải toán có lời văn tốt",
        "comment": ""
    },
    {
        "stt": 3,
        "code": "HS0303",
        "name": "Lê Bảo Nam",
        "gender": "Nam",
        "dob": "04/05/2016",
        "score": 8.5,
        "level": "T",
        "note": "Hiểu bài nhanh, đôi khi còn tính nhẩm ẩu",
        "comment": ""
    },
    {
        "stt": 4,
        "code": "HS0304",
        "name": "Phạm Thục Quyên",
        "gender": "Nữ",
        "dob": "19/11/2016",
        "score": 8.0,
        "level": "H",
        "note": "Ngoan, chăm chỉ, hoàn thành bài tốt",
        "comment": ""
    },
    {
        "stt": 5,
        "code": "HS0305",
        "name": "Vũ Đăng Khoa",
        "gender": "Nam",
        "dob": "02/02/2016",
        "score": 7.5,
        "level": "H",
        "note": "Biết cách giải bài nhưng còn lúng túng phần vẽ hình",
        "comment": ""
    },
    {
        "stt": 6,
        "code": "HS0306",
        "name": "Hoàng Thùy Chi",
        "gender": "Nữ",
        "dob": "14/09/2016",
        "score": 7.0,
        "level": "H",
        "note": "Làm tính tốt, cần đọc kỹ đề toán có lời văn",
        "comment": ""
    },
    {
        "stt": 7,
        "code": "HS0307",
        "name": "Đỗ Gia Huy",
        "gender": "Nam",
        "dob": "30/07/2016",
        "score": 6.5,
        "level": "H",
        "note": "Hay quên nhớ trong phép nhân chia, cần rèn luyện thêm",
        "comment": ""
    },
    {
        "stt": 8,
        "code": "HS0308",
        "name": "Bùi Tuệ Mẫn",
        "gender": "Nữ",
        "dob": "08/10/2016",
        "score": 6.0,
        "level": "H",
        "note": "Tiếp thu bài có tiến bộ, chữ viết cần nắn nót hơn",
        "comment": ""
    },
    {
        "stt": 9,
        "code": "HS0309",
        "name": "Đinh Quốc Bảo",
        "gender": "Nam",
        "dob": "17/04/2016",
        "score": 4.5,
        "level": "C",
        "note": "Còn lúng túng khi làm toán chia, cần giáo viên kèm cặp thêm",
        "comment": ""
    },
    {
        "stt": 10,
        "code": "HS0310",
        "name": "Trịnh Khánh Linh",
        "gender": "Nữ",
        "dob": "22/12/2016",
        "score": 9.0,
        "level": "T",
        "note": "Giải toán nhanh nhẹn, trình bày bài sạch sẽ",
        "comment": ""
    },
    {
        "stt": 11,
        "code": "HS0311",
        "name": "Ngô Tuấn Kiệt",
        "gender": "Nam",
        "dob": "05/06/2016",
        "score": 5.5,
        "level": "H",
        "note": "Cần tập trung hơn trong giờ học, tính toán còn chậm",
        "comment": ""
    },
    {
        "stt": 12,
        "code": "HS0312",
        "name": "Dương Thảo Vy",
        "gender": "Nữ",
        "dob": "11/01/2016",
        "score": 8.5,
        "level": "T",
        "note": "Chăm ngoan, phát biểu sôi nổi, cẩn thận khi làm bài",
        "comment": ""
    },
    {
        "stt": 13,
        "code": "HS0313",
        "name": "Lâm Quốc Cường",
        "gender": "Nam",
        "dob": "29/03/2016",
        "score": 4.0,
        "level": "C",
        "note": "Chưa thuộc bảng nhân chia, cần phụ đạo thêm giờ ra chơi",
        "comment": ""
    },
    {
        "stt": 14,
        "code": "HS0314",
        "name": "Mai Thục Anh",
        "gender": "Nữ",
        "dob": "15/10/2016",
        "score": 10.0,
        "level": "T",
        "note": "Xuất sắc, tư duy sáng tạo, giúp đỡ các bạn trong nhóm",
        "comment": ""
    },
    {
        "stt": 15,
        "code": "HS0315",
        "name": "Hồ Nhật Nam",
        "gender": "Nam",
        "dob": "09/07/2016",
        "score": 7.0,
        "level": "H",
        "note": "Có tiến bộ về tính toán, cần chú ý ghi đúng đơn vị đo",
        "comment": ""
    }
]


def create_sample_excel_workbook() -> Workbook:
    """Creates a beautifully styled standard primary gradebook workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Theo_Doi_Danh_Gia_TT27"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY_HEADER = "1B3B6F"
    LIGHT_BG = "F4F7FB"
    ZEBRA_BG = "F9FBFC"
    BORDER_COLOR = "D0D7DE"

    font_title = Font(name="Segoe UI", size=14, bold=True, color="1B3B6F")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="4A5568")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="1A202C")
    font_bold_data = Font(name="Segoe UI", size=10, bold=True, color="1A202C")

    fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR),
    )

    # Title rows
    ws.merge_cells("A1:I1")
    ws["A1"] = "BẢNG THEO DÕI VÀ ĐÁNH GIÁ HỌC SINH THEO THÔNG TƯ 27/2020/TT-BGDĐT"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = "Trường Tiểu Học Ánh Dương | Lớp: 3A | Môn: Toán | Đánh giá: Cuối Học kỳ 1 | Năm học: 2024-2025"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Spacer
    ws.row_dimensions[3].height = 10

    # Header row at row 4
    headers = [
        ("STT", 6),
        ("Mã HS", 12),
        ("Họ và tên", 26),
        ("Giới tính", 10),
        ("Ngày sinh", 14),
        ("Điểm số", 10),
        ("Mức đạt (T/H/C)", 15),
        ("Ghi chú giáo viên", 32),
        ("Nhận xét Thông tư 27", 45),
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, (header_name, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # Data rows starting from row 5
    for row_idx, item in enumerate(SAMPLE_STUDENTS, start=5):
        ws.row_dimensions[row_idx].height = 22
        current_fill = fill_zebra if row_idx % 2 == 0 else fill_white

        row_values = [
            (item["stt"], Alignment(horizontal="center", vertical="center"), font_data),
            (item["code"], Alignment(horizontal="center", vertical="center"), font_data),
            (item["name"], Alignment(horizontal="left", vertical="center"), font_bold_data),
            (item["gender"], Alignment(horizontal="center", vertical="center"), font_data),
            (item["dob"], Alignment(horizontal="center", vertical="center"), font_data),
            (item["score"], Alignment(horizontal="center", vertical="center"), font_bold_data),
            (item["level"], Alignment(horizontal="center", vertical="center"), font_bold_data),
            (item["note"], Alignment(horizontal="left", vertical="center"), font_data),
            (item["comment"], Alignment(horizontal="left", vertical="center", wrap_text=True), font_data),
        ]

        for col_idx, (val, align, font_style) in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_style
            cell.alignment = align
            cell.fill = current_fill
            cell.border = thin_border

    return wb


SAMPLE_STUDENTS_IMG = [
    ("Trần Thiên", "An", "T", 9.0),
    ("Nguyễn Mai", "Anh", "T", 10.0),
    ("Trương Mai", "Anh", "H", 8.0),
    ("Nguyễn Huỳnh Gia", "Bảo", "T", 9.5),
    ("Trần Nguyễn Ngọc", "Bội", "H", 7.5),
    ("Thạch Huỳnh An", "Châu", "H", 7.0),
    ("Lâm Trần Quốc", "Công", "T", 9.0),
    ("Thạch Kim", "Hoài", "C", 4.5),
    ("Chung Gia", "Hưng", "H", 7.5),
    ("Phạm Gia", "Khang", "T", 9.0),
    ("Trần Tuấn", "Khang", "H", 8.0),
    ("Thạch Ngọc Thiên", "Kim", "T", 9.5),
    ("Trương Thị Kiều", "Loan", "H", 7.0),
    ("Huỳnh Bảo", "Long", "T", 8.5),
    ("Thạch Minh", "Luân", "C", 4.0),
]


def create_sample_midterm_excel_workbook() -> Workbook:
    """Creates a sample workbook exactly matching Image 1 (Giữa Kỳ 1 - Lớp 1 - Môn Tiếng Việt)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tiếng Việt"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=13, bold=True)
    font_sub = Font(name="Times New Roman", size=12, bold=True)
    font_italic = Font(name="Times New Roman", size=11, italic=True)
    font_header = Font(name="Times New Roman", size=11, bold=True)
    font_data = Font(name="Times New Roman", size=11)

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    ws["A1"] = "ỦY BAN NHÂN DÂN XÃ NGỌC TỐ"
    ws["A1"].font = font_sub
    ws["E1"] = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    ws["E1"].font = font_sub
    ws["E1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "TRƯỜNG TIỂU HỌC NGỌC TỐ"
    ws["A2"].font = font_sub
    ws["E2"] = "Độc lập - Tự do - Hạnh phúc"
    ws["E2"].font = font_sub
    ws["E2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A4:F4")
    ws["A4"] = "BẢNG ĐIỂM CHI TIẾT - MÔN TIẾNG VIỆT - HỌC KỲ 1 - GIỮA KỲ 1 - NĂM HỌC 2026 - 2027"
    ws["A4"].font = font_title
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A5:F5")
    ws["A5"] = "Khối 1 - Lớp 1H1"
    ws["A5"].font = font_sub
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    # Header Row at Row 7
    ws["A7"] = "STT"
    ws["A7"].font = font_header
    ws["A7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A7"].border = thin_border

    ws["C7"] = "Họ và tên"
    ws.merge_cells("C7:D7")
    ws["C7"].font = font_header
    ws["C7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["C7"].border = thin_border
    ws["D7"].border = thin_border

    ws["E7"] = "Nhận xét"
    ws["E7"].font = font_header
    ws["E7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E7"].border = thin_border

    ws["F7"] = "XL GK1"
    ws["F7"].font = font_header
    ws["F7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["F7"].border = thin_border

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 3
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["F"].width = 12

    # Data Rows
    for idx, (ho_dem, ten, level, _) in enumerate(SAMPLE_STUDENTS_IMG, start=1):
        r = 7 + idx
        ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=3, value=ho_dem).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=4, value=ten).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=5, value="").alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=6, value=level).alignment = Alignment(horizontal="center", vertical="center")

        for c in [1, 3, 4, 5, 6]:
            ws.cell(row=r, column=c).border = thin_border
            ws.cell(row=r, column=c).font = font_data

    return wb


def create_sample_endterm_excel_workbook() -> Workbook:
    """Creates a sample workbook exactly matching Image 2 (Cuối Kỳ 1 - Lớp 1 - Môn Tiếng Việt với cột điểm số KT CK1)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tiếng Việt"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Times New Roman", size=13, bold=True)
    font_sub = Font(name="Times New Roman", size=12, bold=True)
    font_header = Font(name="Times New Roman", size=11, bold=True)
    font_data = Font(name="Times New Roman", size=11)

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    ws["A1"] = "ỦY BAN NHÂN DÂN XÃ NGỌC TỐ"
    ws["A1"].font = font_sub
    ws["E1"] = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    ws["E1"].font = font_sub
    ws["E1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "TRƯỜNG TIỂU HỌC NGỌC TỐ"
    ws["A2"].font = font_sub
    ws["E2"] = "Độc lập - Tự do - Hạnh phúc"
    ws["E2"].font = font_sub
    ws["E2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A4:G4")
    ws["A4"] = "BẢNG ĐIỂM CHI TIẾT - MÔN TIẾNG VIỆT - HỌC KỲ 1 - CUỐI KỲ 1 - NĂM HỌC 2026 - 2027"
    ws["A4"].font = font_title
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A5:G5")
    ws["A5"] = "Khối 1 - Lớp 1H1"
    ws["A5"].font = font_sub
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    # Header Row at Row 7
    ws["A7"] = "STT"
    ws["A7"].font = font_header
    ws["A7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A7"].border = thin_border

    ws["C7"] = "Họ và tên"
    ws.merge_cells("C7:D7")
    ws["C7"].font = font_header
    ws["C7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["C7"].border = thin_border
    ws["D7"].border = thin_border

    ws["E7"] = "Nhận xét"
    ws["E7"].font = font_header
    ws["E7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E7"].border = thin_border

    ws["F7"] = "KT CK1"
    ws["F7"].font = font_header
    ws["F7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["F7"].border = thin_border

    ws["G7"] = "XL CK1"
    ws["G7"].font = font_header
    ws["G7"].alignment = Alignment(horizontal="center", vertical="center")
    ws["G7"].border = thin_border

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 3
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10

    # Data Rows
    for idx, (ho_dem, ten, level, score) in enumerate(SAMPLE_STUDENTS_IMG, start=1):
        r = 7 + idx
        ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=3, value=ho_dem).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=4, value=ten).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=5, value="").alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=6, value=score).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=7, value=level).alignment = Alignment(horizontal="center", vertical="center")

        for c in [1, 3, 4, 5, 6, 7]:
            ws.cell(row=r, column=c).border = thin_border
            ws.cell(row=r, column=c).font = font_data

    return wb


def get_sample_midterm_excel_bytes() -> bytes:
    """Returns sample Giữa Kỳ 1 Excel workbook matching Image 1."""
    wb = create_sample_midterm_excel_workbook()
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def get_sample_endterm_excel_bytes() -> bytes:
    """Returns sample Cuối Kỳ 1 Excel workbook with KT CK1 matching Image 2."""
    wb = create_sample_endterm_excel_workbook()
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def get_sample_excel_bytes() -> bytes:
    """Returns standard sample Excel workbook as bytes."""
    return get_sample_midterm_excel_bytes()



def get_sample_exam_image_bytes() -> bytes:
    """Generates a realistic primary school exam image (PNG) for instant OCR testing."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 1150), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    # Header
    draw.text((40, 30), "PHÒNG GD&ĐT QUẬN HOÀN KIẾM", fill=(30, 41, 59))
    draw.text((40, 50), "TRƯỜNG TIỂU HỌC KIM ĐỒNG", fill=(30, 41, 59))

    draw.text((530, 30), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", fill=(30, 41, 59))
    draw.text((580, 50), "Độc lập - Tự do - Hạnh phúc", fill=(30, 41, 59))
    draw.line([(580, 70), (770, 70)], fill=(148, 163, 184), width=1)

    # Title
    draw.text((250, 100), "ĐỀ KIỂM TRA ĐỊNH KỲ CUỐI HỌC KỲ I", fill=(30, 58, 138))
    draw.text((290, 125), "NĂM HỌC 2025 - 2026 • MÔN: TOÁN LỚP 3", fill=(15, 23, 42))
    draw.text((360, 145), "(Thời gian làm bài: 40 phút)", fill=(100, 116, 139))

    # Student info box
    draw.rectangle([(40, 175), (860, 235)], outline=(203, 213, 225), fill=(248, 250, 252), width=1)
    draw.text((55, 190), "Họ và tên: Nguyễn Hoàng Minh Khôi                         Lớp: 3A1            Điểm: ..........", fill=(15, 23, 42))
    draw.text((55, 210), "Lời nhận xét của giáo viên: .....................................................................................", fill=(100, 116, 139))

    # Part 1: Multiple choice
    draw.text((40, 255), "I. PHẦN TRẮC NGHIỆM (4.0 điểm)", fill=(30, 58, 138))
    draw.text((40, 280), "Khoanh tròn vào chữ cái đặt trước câu trả lời đúng nhất:", fill=(51, 65, 85))

    draw.text((40, 310), "Câu 1 (1.0 điểm): Số lớn nhất có ba chữ số khác nhau là:", fill=(15, 23, 42))
    draw.text((70, 330), "A. 999               B. 987               C. 989               D. 978", fill=(30, 41, 59))

    draw.text((40, 365), "Câu 2 (1.0 điểm): Một hình vuông có cạnh dài 8cm. Chu vi của hình vuông đó là:", fill=(15, 23, 42))
    draw.text((70, 385), "A. 16 cm             B. 24 cm             C. 32 cm             D. 64 cm", fill=(30, 41, 59))

    draw.text((40, 420), "Câu 3 (1.0 điểm): Đã tô màu 1/3 hình nào dưới đây?", fill=(15, 23, 42))
    draw.text((70, 440), "A. Hình 1            B. Hình 2            C. Hình 3            D. Hình 4", fill=(30, 41, 59))

    draw.text((40, 475), "Câu 4 (1.0 điểm): Tìm x biết:  x : 6 = 14", fill=(15, 23, 42))
    draw.text((70, 495), "A. x = 84            B. x = 20            C. x = 72            D. x = 96", fill=(30, 41, 59))

    # Part 2: Essay & Calculations
    draw.text((40, 540), "II. PHẦN TỰ LUẬN (6.0 điểm)", fill=(30, 58, 138))

    draw.text((40, 570), "Bài 1 (2.0 điểm): Đặt tính rồi tính:", fill=(15, 23, 42))
    draw.text((70, 590), "a) 234 + 185                 b) 652 - 279                 c) 124 x 3                 d) 84 : 4", fill=(30, 41, 59))

    draw.text((40, 635), "Bài 2 (2.0 điểm): Một cửa hàng có 72 gói bánh. Cửa hàng đã bán được 1/4 số gói bánh đó.", fill=(15, 23, 42))
    draw.text((40, 655), "Hỏi cửa hàng còn lại bao nhiêu gói bánh?", fill=(15, 23, 42))

    # Table in Exam
    draw.text((40, 705), "Bài 3 (2.0 điểm): Hoàn thành bảng thống kê số cây lớp 3A đã trồng:", fill=(15, 23, 42))

    # Table header
    draw.rectangle([(40, 735), (860, 770)], fill=(224, 231, 255), outline=(148, 163, 184))
    draw.text((60, 745), "Tổ / Phân nhóm", fill=(30, 58, 138))
    draw.text((260, 745), "Loại cây trồng", fill=(30, 58, 138))
    draw.text((480, 745), "Số lượng (cây)", fill=(30, 58, 138))
    draw.text((680, 745), "Đánh giá kết quả", fill=(30, 58, 138))

    rows = [
        ("Tổ 1", "Cây hoa cúc", "25", "Đạt chỉ tiêu"),
        ("Tổ 2", "Cây hoa hồng", "30", "Vượt chỉ tiêu"),
        ("Tổ 3", "Cây hoa mười giờ", "28", "Đạt chỉ tiêu"),
        ("Tổng cộng", "Tất cả các loại hoa", "83 cây", "Xuất sắc phong trào"),
    ]
    y = 770
    for r in rows:
        draw.rectangle([(40, y), (860, y + 32)], fill=(255, 255, 255), outline=(203, 213, 225))
        draw.text((60, y + 8), r[0], fill=(15, 23, 42))
        draw.text((260, y + 8), r[1], fill=(15, 23, 42))
        draw.text((480, y + 8), r[2], fill=(15, 23, 42))
        draw.text((680, y + 8), r[3], fill=(15, 23, 42))
        y += 32

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def get_sample_exam_pdf_bytes() -> bytes:
    """Generates a realistic primary school exam PDF for instant OCR testing."""
    from PIL import Image
    png_bytes = get_sample_exam_image_bytes()
    img = Image.open(io.BytesIO(png_bytes))
    pdf_buf = io.BytesIO()
    img.convert("RGB").save(pdf_buf, format="PDF")
    return pdf_buf.getvalue()
